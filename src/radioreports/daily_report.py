from __future__ import annotations

from pathlib import Path
from datetime import datetime
import os
from typing import Optional

import logging
from .logging_cfg import setup_logging

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import file_utils, docx_utils, report_utils
from .settings import AREAS, TB_CONFIG, FOOTER_SIGN, FOOTER_TEXT_1, FOOTER_TEXT_2
from . import settings

FILE_MASK: str = "report_*.xlsx"


# ---------- Styles helpers ----------
def ensure_para_style(doc: Document, name: str, *, size_pt: int = 14, font_name: str = "Times New Roman"):
    """Create or return a PARAGRAPH style with given formatting."""
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(size_pt)
    style.font.name = font_name
    return style


def ensure_char_style(doc: Document, name: str, *, size_pt: int = 14, font_name: str = "Times New Roman"):
    """Create or return a CHARACTER style with given formatting."""
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    style.font.size = Pt(size_pt)
    style.font.name = font_name
    return style


# ---------- Main ----------
def main() -> None:
    
    setup_logging(settings.LOG_LEVEL, settings.OUTPUT_DIR)                     # ← додай
    log = logging.getLogger(__name__)   # ← додай

    log.info("Старт генерації щоденного звіту")
    
    # --- find latest XLSX ---
    try:
        xlsx_path = Path(file_utils.get_fresh_communitify_report(FILE_MASK))
    except FileNotFoundError as e:
        log.error(str(e))
        print("\n⚠️  Перевір DATA_DIR у .env та наявність файлів за маскою:", FILE_MASK)
        input("Натисни Enter, щоб закрити...")   # щоб вікно не закрилось одразу
        return
    
    file_name_only = xlsx_path.name

    today = datetime.now().strftime("%d.%m.%Y")
    start_hour = report_utils.get_start_hour(file_name_only)
    end_hour = report_utils.get_time_end(file_name_only)[:2]
    message = f"Create report for {file_name_only} from {start_hour} to {end_hour}"
    log.info(message)

    is_hur_report = start_hour in {"11", "23"}

    # --- load and transform data ---
    df = file_utils.get_df_from_xlsx(str(xlsx_path))

    # Helper: normalize headers for robust matching (tolerant to spaces and \\ vs /)
    def _norm(s: str) -> str:
        return (
            str(s)
            .strip()
            .lower()
            .replace(" ", "")
            .replace("\\", "/")
        )

    cols = list(df.columns)
    norm_map = {c: _norm(c) for c in cols}

    def find_col(*candidates: str) -> Optional[str]:
        cand_norm = [_norm(c) for c in candidates]
        for original, n in norm_map.items():
            if any(n == cn or cn in n for cn in cand_norm):
                return original
        return None

    # detect key column names from various possible spellings
    col_notes = find_col("Висновки", "примітки")
    col_radio = find_col("р/обмін", "р\\обмін", "радіоперехоплення")
    col_loc   = find_col("Локація", "координати")
    col_date  = find_col("Дата")
    col_time  = find_col("Час")
    col_freq  = find_col("Частота")
    col_name  = find_col("Назвар/м", "Назвар\\м", "Назва р/м", "Назва р\\м")
    col_who   = find_col("хто")
    col_to    = find_col("кому")

    # sort if date/time exist
    sort_cols = [c for c in [col_date, col_time] if c]
    sorted_df = df.sort_values(by=sort_cols) if sort_cols else df.copy()

    # rename into canonical names if present
    rename_map = {}
    if col_notes: rename_map[col_notes] = "Висновки"
    if col_radio: rename_map[col_radio] = "Радіоперехоплення"
    if col_loc:   rename_map[col_loc]   = "Локація"
    sorted_df = sorted_df.rename(columns=rename_map)

    # drop auxiliary columns if they exist
    drop_cols = [c for c in [col_date, col_time, col_freq, col_name, col_who, col_to] if c]
    if drop_cols:
        sorted_df = sorted_df.drop(drop_cols, axis=1, errors="ignore")

    # reorder columns if all present, else take intersection
    desired = ["Висновки", "Радіоперехоплення", "Локація"]
    present = [c for c in desired if c in sorted_df.columns]
    df_reordered = sorted_df.loc[:, present].astype(str)  # ensure strings

    # set area id into "Локація" based on text matches in "Висновки"
    if "Локація" in df_reordered.columns and "Висновки" in df_reordered.columns:
        for idx in df_reordered.index:
            description: str = str(df_reordered.loc[idx, "Висновки"])  # ensure string for len()
            if len(description) < 5:
                continue
            df_reordered.loc[idx, "Локація"] = "0"
            for area in AREAS:
                area_id = area["id"]
                for loc in area["names"]:
                    if loc in description:
                        df_reordered.loc[idx, "Локація"] = str(area_id)
                        break

    # --- build DOCX ---
    doc = Document()

    # header block for HUR report
    if is_hur_report:
        para = doc.add_paragraph("Форма №1.2.17/ОСУВ", style="Normal")
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para.paragraph_format.space_before = Inches(0.25)
        para.paragraph_format.space_after = Inches(0.25)

    # header title (paragraph style + run)
    ensure_para_style(doc, "HeaderPara", size_pt=14, font_name="Times New Roman")

    title = (
        "Донесення\n"
        "за результатами ведення радіоелектронної розвідки\n"
        "у зоні відповідальності тактичної групи “Кремінна”"
    )
    p = doc.add_paragraph(style="HeaderPara")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True

    if is_hur_report:
        period_desc = f"(станом на {int(end_hour) + 1:02d}:00 {today})"
    else:
        period_desc = report_utils.get_description_time(file_name_only)

    p = doc.add_paragraph(period_desc, style="Normal")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # tables per area
    for area in AREAS:
        area_id = area["id"]
        description = area["description"]

        df_area = df_reordered[df_reordered["Локація"] == str(area_id)]

        p = doc.add_paragraph()
        p.add_run(description).bold = True
        p.paragraph_format.space_before = Inches(0.25)
        p.paragraph_format.space_after = Inches(0.25)

        if len(df_area) == 0:
            p = doc.add_paragraph("Не виявлено", style="Normal")
            p.paragraph_format.left_indent = Inches(0.5)
        else:
            df_area = df_area.drop(["Локація"], axis=1)
            docx_utils.add_df_to_docx(doc, df_area, TB_CONFIG)

    if is_hur_report:
        p = doc.add_paragraph(FOOTER_TEXT_1, style="Quote")
        p.paragraph_format.space_before = Inches(0.25)
        doc.add_paragraph(FOOTER_TEXT_2, style="Quote")

    p = doc.add_paragraph(FOOTER_SIGN, style="Normal")
    p.paragraph_format.space_before = Inches(0.5)

    # --- save file to output/<folder> ---
    root = Path(__file__).resolve().parent.parent  # repo root (where src/ lives)
    folder_name = report_utils.get_foler_name_part(file_name_only)  # keep original util name
    out_dir = root / "output" / folder_name
    out_dir.mkdir(parents=True, exist_ok=True)

    base_name = "Звіт РЕР " + ("1.2.17" if is_hur_report else "")
    doc_file = out_dir / f"{base_name} ({folder_name}).docx"

    # архівуємо старі файли у цій теці та зберігаємо новий
    report_utils.move_files_to_archive(str(out_dir), FILE_MASK)

    try:
        doc.save(str(doc_file))
    except Exception as e:
        print(f"Error saving document: {e}")
    else:
        print(f"✅ Work is done! {doc_file}")
        # відкрити теку та файл (Windows only)
        try:
            os.startfile(str(out_dir))
            os.startfile(str(doc_file))
        except Exception:
            pass
        
    log.info("Звіт успішно створено")



if __name__ == "__main__":
    main()
