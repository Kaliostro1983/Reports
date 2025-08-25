# src/radioreports/active_networks_63.py
from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable, Optional

import logging
import os
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from zoneinfo import ZoneInfo

from . import settings
from .io.freqs_loader import load_freqs
from .logging_cfg import setup_logging

# ---- очікувані назви колонок у довіднику ----
COL_FREQ = "Частота"
# можливі назви колонки з описом/назвою мережі
NET_NAME_CANDIDATES = ["Радіомережа", "Назва р/м", "Назвар/м", "Назва", "Назва_мережі", "Опис частоти"]
COL_STATUS = "Статус"
ACTIVE_VALUE = "Спостерігається"

FOOTER_NOTE = "Решта частот — ймовірно закрита цифра або похибка пеленгації/ідентифікації, " "що потребує додаткової перевірки."

_TZ = ZoneInfo("Europe/Kyiv")


# ---------------- helpers ----------------
def _fmt_freq4(val) -> str:
    """4 знаки після крапки або порожньо."""
    if pd.isna(val):
        return ""
    try:
        return f"{float(val):.4f}"
    except Exception:
        return str(val)


def _set_col_widths(table, widths_in_inches: list[float]) -> None:
    """Жорстко задає ширину колонок для всіх рядків (inch -> twips)."""
    for row in table.rows:
        for idx, width in enumerate(widths_in_inches):
            tc = row.cells[idx]._tc
            tcPr = tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcW")):
                tcPr.remove(el)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcPr.append(tcW)


def _fit_picture_to_page(doc: Document, picture) -> None:
    """Пропорційно підганяє зображення під область друку сторінки."""
    section = doc.sections[0]
    max_w = section.page_width - section.left_margin - section.right_margin
    max_h = section.page_height - section.top_margin - section.bottom_margin
    w, h = picture.width, picture.height
    scale = min(max_w / w, max_h / h, 1.0)
    if scale < 1.0:
        picture.width = int(w * scale)
        picture.height = int(h * scale)


def _find_first_existing(candidates: Iterable[Path]) -> Optional[Path]:
    for p in candidates:
        try:
            if p and Path(p).exists():
                return Path(p)
        except Exception:
            continue
    return None


def _resolve_header_image() -> Optional[Path]:
    """
    Пріоритет зображення:
      1) HEADER_IMAGE (абсолютний або ім'я файлу в IMAGES_DIR)
      2) у IMAGES_DIR — файли з 'header' у назві
      3) у DATA_DIR\\images — файли з 'header' у назві
    """
    images_dir = getattr(settings, "IMAGES_DIR", None)
    header_name = getattr(settings, "HEADER_IMAGE", None)
    data_dir = getattr(settings, "DATA_DIR", None)

    candidates: list[Path] = []

    # 1) HEADER_IMAGE напряму або всередині IMAGES_DIR
    if header_name:
        p = Path(header_name)
        if not p.is_absolute() and images_dir:
            candidates.append(Path(images_dir) / header_name)
        candidates.append(p)

    # 2) будь-які "header*" у IMAGES_DIR
    if images_dir:
        img_dir = Path(images_dir)
        candidates.extend(list(img_dir.glob("header.*")))
        candidates.extend(list(img_dir.glob("header_*.*")))
        candidates.extend(list(img_dir.glob("*header*.*")))

    # 3) DATA_DIR/images — запасний варіант
    if data_dir:
        dd = Path(data_dir) / "images"
        candidates.extend(list(dd.glob("header.*")))
        candidates.extend(list(dd.glob("header_*.*")))
        candidates.extend(list(dd.glob("*header*.*")))

    return _find_first_existing(candidates)


def _resolve_subheader_text() -> Optional[str]:
    """
    Текст під зображенням (підзаголовок).
    Беремо з однієї з змінних, що є у settings/.env:
      SUBHEADER_TITLE | HEADER_TITLE | AREA_TITLE | REGION_TITLE | LOCATION_TITLE
    Якщо нічого не задано — підзаголовок не додаємо.
    """
    for key in ("SUBHEADER_TITLE", "HEADER_TITLE", "AREA_TITLE", "REGION_TITLE", "LOCATION_TITLE"):
        if getattr(settings, key, None):
            return str(getattr(settings, key))
    return None


def _title(doc: Document, text: str) -> None:
    # стиль за замовчуванням
    try:
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
    except Exception:
        pass

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)


def _subtitle(doc: Document, text: str) -> None:
    """Підзаголовок під зображенням — центр, жирний, all caps."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    try:
        run.font.all_caps = True  # візуально як на твоєму прикладі
    except Exception:
        pass
    doc.add_paragraph("")  # невеликий відступ після підзаголовка


def _add_period_info(doc: Document, rows_count: int) -> None:
    """
    Додає текст на кшталт:
    'В поточний період (з 16:00 DD.MM.YYYY по 16:00 DD.MM.YYYY року) відмічено функціонування N р/м:'
    Годину беремо зі settings.PERIOD_START_HOUR (за замовчуванням 16).
    """
    hour = int(getattr(settings, "PERIOD_START_HOUR", 16))
    now = datetime.now(_TZ)
    today = now.date()
    yesterday = today - timedelta(days=1)

    y_s = f"{hour:02d}:00 {yesterday.strftime('%d.%m.%Y')}"
    t_s = f"{hour:02d}:00 {today.strftime('%d.%m.%Y')}"

    p = doc.add_paragraph(f"В поточний період (з {y_s} по {t_s} року) відмічено функціонування {rows_count} р/м:")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _add_table(doc: Document, df: pd.DataFrame, net_col: str) -> None:
    """Рівно 3 колонки: № п/п | Частота | Назва (із df[net_col])."""
    if df.empty:
        doc.add_paragraph("Активних мереж не виявлено за поточними даними.")
        return

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    hdr[0].text = "№ п/п"
    hdr[1].text = "Частота"
    hdr[2].text = "Назва"

    for i, row in df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(i + 1)
        cells[1].text = _fmt_freq4(row[COL_FREQ])
        name = row[net_col]
        cells[2].text = "" if pd.isna(name) else str(name)

    # пропорції: номер вузький, частота компактно, назва — решта
    _set_col_widths(table, [0.6, 1.4, 4.5])


def _footer(doc: Document) -> None:
    doc.add_paragraph("")  # відступ
    p = doc.add_paragraph(FOOTER_NOTE)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sign = getattr(settings, "FOOTER_SIGN", "")
    if sign:
        doc.add_paragraph("")
        p2 = doc.add_paragraph(sign)
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if p2.runs:
            p2.runs[0].font.size = Pt(10)
            p2.runs[0].italic = True


# ---------------- main ----------------
def main() -> int:
    setup_logging(settings.LOG_LEVEL, settings.OUTPUT_DIR)
    log = logging.getLogger(__name__)

    # 1) Дані: готовий лоадер частот
    try:
        freqs = load_freqs()
    except Exception as e:
        log.exception("Не вдалося завантажити довідник частот: %s", e)
        print("❌ Помилка завантаження довідника частот. Перевір FREQS_FILE у .env")
        return 1

    # 2) Відфільтрувати тільки активні
    df = freqs.copy()
    if COL_STATUS not in df.columns:
        raise KeyError(f"У файлі немає колонки '{COL_STATUS}'")
    df = df[df[COL_STATUS].astype(str).str.strip() == ACTIVE_VALUE]

    # 3) Визначити колонку з назвою мережі
    net_col = next((c for c in NET_NAME_CANDIDATES if c in df.columns), None)
    if not net_col:
        raise KeyError("У файлі не знайдено колонку з назвою мережі. Очікувались один із заголовків: " + ", ".join(NET_NAME_CANDIDATES))

    # 4) Відібрати потрібні поля та відсортувати за частотою
    if COL_FREQ not in df.columns:
        raise KeyError(f"У файлі немає колонки '{COL_FREQ}'")

    df = df[[COL_FREQ, net_col]].dropna(subset=[COL_FREQ]).sort_values(COL_FREQ, kind="mergesort").reset_index(drop=True)

    # 5) DOCX
    doc = Document()

    today = datetime.now(_TZ).strftime("%d.%m.%Y")
    _title(doc, f"Активні мережі (63 омсбр)\nстаном на {today}")

    header_img = _resolve_header_image()
    if header_img:
        try:
            pic = doc.add_picture(str(header_img))
            _fit_picture_to_page(doc, pic)
            doc.add_paragraph("")  # відступ після картинки
        except Exception as e:
            log.warning("Не вдалося вставити зображення %s: %s", header_img, e)
    else:
        log.info("HEADER_IMAGE не знайдено (IMAGES_DIR/HEADER_IMAGE) — пропускаємо зображення.")

    # --- новий підзаголовок під зображенням (якщо задано в settings/.env) ---
    sub = _resolve_subheader_text()
    if sub:
        _subtitle(doc, sub)

    # --- службовий текст і таблиця ---
    _add_period_info(doc, len(df))
    doc.add_paragraph("")  # маленький відступ перед таблицею

    _add_table(doc, df, net_col)
    _footer(doc)

    # 6) Збереження
    out_dir = Path(getattr(settings, "OUTPUT_DIR", Path.cwd() / "output"))
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"Активні мережі (63 омсбр) {today}.docx"

    # якщо файл відкритий — з суфіксом
    for i in range(20):
        try:
            doc.save(str(out_path))
            break
        except Exception:
            i += 1
            out_path = out_dir / f"Активні мережі (63 омсбр) {today}_в{i}.docx"

    log.info("Готово: %s", out_path)
    print(f"✅ Звіт сформовано: {out_path}")

    try:
        os.startfile(str(out_path))
        os.startfile(str(out_dir))
    except Exception:
        pass

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
