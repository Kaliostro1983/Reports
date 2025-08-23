from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from pandas import DataFrame


def add_df_to_docx(doc: Document, df: DataFrame, tb_config: list):
    df = df
    config = tb_config
    headers = df.columns
    num_of_rows = len(df)
    num_of_columns = len(headers)

    table = doc.add_table(rows=1, cols=num_of_columns)
    table.style = "Table Grid"

    # add headers
    row = table.rows[0].cells
    for i in range(len(headers)):
        cell = row[i]
        cell.text = headers[i]
        cell.width = Inches(config[i]["width_inch"])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # fill table by data
    for i in range(num_of_rows):
        row = table.add_row().cells
        for j in range(len(headers)):
            cell = row[j]
            cell.text = str(df.iat[i, j])
            cell.width = Inches(config[j]["width_inch"])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.alignment = config[j]["h_align"]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
