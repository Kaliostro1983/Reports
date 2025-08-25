# -*- coding: utf-8 -*-
from __future__ import annotations
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

GREEN = "70AD47"  # Мертва + є перехоплення
RED = "F40000"  # Не мертва + 0 перехоплень


def save_activity_xlsx(df: pd.DataFrame, day_labels, out_path: str) -> str:
    """
    Запис підсумкової таблиці у XLSX:
      - жирний заголовок, автофільтр, freeze panes
      - номерний формат для колонки 'Частота' = 0.0000 (4 знаки після коми)
      - центрування денних колонок
      - підсвітка рядків за правилами
      - унікальний суфікс, якщо файл відкритий
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Активність"

    # Запис DataFrame
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    n_rows = df.shape[0] + 1  # включно з хедером
    n_cols = df.shape[1]

    # Заголовки
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    # Ширини колонок (мінімально розумні)
    widths = {1: 7, 2: 12, 3: 18}
    for col in range(1, n_cols + 1):
        letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 9)

    # --- Формат 'Частота' як NUMBER з 4 знаками після коми ---
    # Колонка 2 = 'Частота'
    for r in range(2, n_rows + 1):
        c = ws.cell(row=r, column=2)
        # Лише якщо це число (не чіпаємо NaN/порожні)
        try:
            float_val = float(c.value)
            c.value = float_val
            c.number_format = "0.0000"
            c.alignment = Alignment(horizontal="right", vertical="center")
        except Exception:
            pass

    # Денні колонки центруємо
    first_day_col = 4  # 1:№, 2:Частота, 3:Статус, з 4-ї йдуть дні
    for col in range(first_day_col, n_cols + 1):
        for r in range(2, n_rows + 1):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", vertical="center")

    # Підсвітка рядків
    for r in range(2, n_rows + 1):
        status = ws.cell(row=r, column=3).value
        total = 0
        for c in range(first_day_col, first_day_col + len(day_labels)):
            try:
                total += int(ws.cell(row=r, column=c).value or 0)
            except Exception:
                pass

        fill = None
        if status == "Мертва" and total > 0:
            fill = PatternFill("solid", fgColor=GREEN)
        elif status != "Мертва" and total == 0:
            fill = PatternFill("solid", fgColor=RED)

        if fill:
            for c in range(1, n_cols + 1):
                ws.cell(row=r, column=c).fill = fill

    # Збереження з унікальним суфіксом, якщо файл відкритий
    try:
        wb.save(out_path)
        return out_path
    except PermissionError:
        base, ext = os.path.splitext(out_path)
        alt = f"{base}__{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        wb.save(alt)
        return alt
